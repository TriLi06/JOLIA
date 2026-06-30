from __future__ import annotations

import logging
from pathlib import Path

from app.processors.base_processor import BaseProcessor, ProcessingResult
from app.services import sidecar_service
from app.services.chunking_service import split_text

logger = logging.getLogger(__name__)


class PptxProcessor(BaseProcessor):
    processor_version = "1.0"

    def process(self, file_path: Path, file_record, config) -> ProcessingResult:
        try:
            from pptx import Presentation

            prs = Presentation(str(file_path))
        except Exception as exc:
            return ProcessingResult(success=False, error_message=f"PPTX-Lesefehler: {exc}")

        slide_texts: list[str] = []
        for i, slide in enumerate(prs.slides, start=1):
            parts: list[str] = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for para in shape.text_frame.paragraphs:
                        text = para.text.strip()
                        if text:
                            parts.append(text)
            if parts:
                slide_texts.append(f"[Folie {i}]\n" + "\n".join(parts))

        full_text = "\n\n".join(slide_texts)

        metadata = {
            "Dateigröße": f"{file_record.file_size or 0:,} Bytes",
            "SHA256": file_record.sha256,
            "Folienanzahl": len(prs.slides),
            "Wortanzahl": len(full_text.split()),
        }

        json_data = {
            "file_id": file_record.id,
            "original_filename": file_record.original_filename,
            "sha256": file_record.sha256,
            "mime_type": file_record.mime_type,
            "content_type": "documents",
            "file_size": file_record.file_size,
            "processor": "PptxProcessor",
            "processor_version": self.processor_version,
            "slide_count": len(prs.slides),
            "word_count": len(full_text.split()),
        }
        json_path = sidecar_service.write_json_sidecar(file_path, json_data)
        md_content = sidecar_service.build_document_md(
            file_record.original_filename, "PowerPoint-Präsentation", full_text, metadata
        )
        md_path = sidecar_service.write_md_sidecar(file_path, md_content)

        chunks = split_text(
            full_text,
            chunk_size=config.processing.chunk_size,
            overlap=config.processing.chunk_overlap,
        )

        return ProcessingResult(
            success=True,
            chunks=chunks,
            metadata=metadata,
            sidecar_json_path=str(json_path),
            sidecar_md_path=str(md_path),
        )
