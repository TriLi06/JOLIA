from __future__ import annotations

import logging
from pathlib import Path

from app.processors.base_processor import BaseProcessor, ProcessingResult
from app.services import sidecar_service
from app.services.chunking_service import split_text

logger = logging.getLogger(__name__)


class DocxProcessor(BaseProcessor):
    processor_version = "1.0"

    def process(self, file_path: Path, file_record, config) -> ProcessingResult:
        try:
            from docx import Document

            doc = Document(str(file_path))
        except Exception as exc:
            return ProcessingResult(success=False, error_message=f"DOCX-Lesefehler: {exc}")

        parts: list[str] = []

        for para in doc.paragraphs:
            if para.text.strip():
                parts.append(para.text.strip())

        for table in doc.tables:
            table_rows: list[str] = []
            headers = [cell.text.strip() for cell in table.rows[0].cells] if table.rows else []
            if headers:
                table_rows.append("| " + " | ".join(headers) + " |")
                table_rows.append("| " + " | ".join(["---"] * len(headers)) + " |")
            for row in table.rows[1:]:
                cells = [cell.text.strip() for cell in row.cells]
                table_rows.append("| " + " | ".join(cells) + " |")
            if table_rows:
                parts.append("\n".join(table_rows))

        full_text = "\n\n".join(parts)

        metadata = {
            "Dateigröße": f"{file_record.file_size or 0:,} Bytes",
            "SHA256": file_record.sha256,
            "Absätze": len(doc.paragraphs),
            "Tabellen": len(doc.tables),
            "Wortanzahl": len(full_text.split()),
        }

        json_data = {
            "file_id": file_record.id,
            "original_filename": file_record.original_filename,
            "sha256": file_record.sha256,
            "mime_type": file_record.mime_type,
            "content_type": "documents",
            "file_size": file_record.file_size,
            "processor": "DocxProcessor",
            "processor_version": self.processor_version,
            "paragraphs": len(doc.paragraphs),
            "tables": len(doc.tables),
            "word_count": len(full_text.split()),
        }
        json_path = sidecar_service.write_json_sidecar(file_path, json_data)
        md_content = sidecar_service.build_document_md(
            file_record.original_filename, "Word-Dokument", full_text, metadata
        )
        md_path = sidecar_service.write_md_sidecar(file_path, md_content)

        chunks = split_text(
            full_text,
            chunk_size=config.processing.chunk_size,
            overlap=config.processing.chunk_overlap,
        )

        return ProcessingResult(
            success=True,
            chunks=chunks,
            metadata=metadata,
            sidecar_json_path=str(json_path),
            sidecar_md_path=str(md_path),
        )
