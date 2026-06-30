from __future__ import annotations

import logging
from pathlib import Path

from app.processors.base_processor import BaseProcessor, ProcessingResult
from app.services import sidecar_service
from app.services.chunking_service import TextChunk, split_text

logger = logging.getLogger(__name__)


class XlsxProcessor(BaseProcessor):
    processor_version = "1.0"

    def process(self, file_path: Path, file_record, config) -> ProcessingResult:
        try:
            from openpyxl import load_workbook

            wb = load_workbook(str(file_path), read_only=True, data_only=True)
        except Exception as exc:
            return ProcessingResult(success=False, error_message=f"XLSX-Lesefehler: {exc}")

        sheet_parts: list[str] = []
        total_rows = 0

        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            rows: list[list[str]] = []
            for row in ws.iter_rows(values_only=True):
                cells = [str(c) if c is not None else "" for c in row]
                if any(c.strip() for c in cells):
                    rows.append(cells)
                    total_rows += 1

            if not rows:
                continue

            md_rows = []
            if rows:
                header = rows[0]
                md_rows.append("| " + " | ".join(header) + " |")
                md_rows.append("| " + " | ".join(["---"] * len(header)) + " |")
                for data_row in rows[1:]:
                    md_rows.append("| " + " | ".join(data_row) + " |")

            sheet_parts.append(f"## Tabellenblatt: {sheet_name}\n\n" + "\n".join(md_rows))

        full_text = "\n\n".join(sheet_parts)

        metadata = {
            "Dateigröße": f"{file_record.file_size or 0:,} Bytes",
            "SHA256": file_record.sha256,
            "Tabellenblätter": len(wb.sheetnames),
            "Zeilen gesamt": total_rows,
        }

        json_data = {
            "file_id": file_record.id,
            "original_filename": file_record.original_filename,
            "sha256": file_record.sha256,
            "mime_type": file_record.mime_type,
            "content_type": "documents",
            "file_size": file_record.file_size,
            "processor": "XlsxProcessor",
            "processor_version": self.processor_version,
            "sheets": wb.sheetnames,
            "total_rows": total_rows,
        }
        json_path = sidecar_service.write_json_sidecar(file_path, json_data)
        md_content = sidecar_service.build_document_md(
            file_record.original_filename, "Excel-Tabelle", full_text, metadata
        )
        md_path = sidecar_service.write_md_sidecar(file_path, md_content)

        chunks = split_text(
            full_text,
            chunk_size=config.processing.chunk_size,
            overlap=config.processing.chunk_overlap,
        )

        return ProcessingResult(
            success=True,
            chunks=chunks,
            metadata=metadata,
            sidecar_json_path=str(json_path),
            sidecar_md_path=str(md_path),
        )
